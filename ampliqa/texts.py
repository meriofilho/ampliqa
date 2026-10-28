"""Rascunhos de Métodos, Análise estatística, Resultados e Legenda (inglês e português), preenchidos com a análise.

O padrão segue os POPs atuais (PO.OPE.01 a 04BIOMOL); o usuário confirma kits e equipamentos na tela Textos.
Os textos são rascunhos: o pesquisador deve revisar antes de usar.
"""
import io

import numpy as np
import pandas as pd

from .design import TISSUES
from .stats import stars

# tipo de amostra (rótulo pt) -> (frase en, frase pt) no plural; nomes antigos continuam valendo
SAMPLE_TYPES = {k: (v[1], v[2]) for k, v in TISSUES.items()}
_OLD = {"Larvas inteiras": "Larva inteira", "Cérebros": "Cérebro", "Cabeças": "Cabeça", "Fígados": "Fígado"}

METHODS_DEFAULTS = {
    "especie": "zebrafish (Danio rerio)",
    "especie_pt": "zebrafish (Danio rerio)",
    "tipo_amostra": "Larva inteira",
    "animais_por_pool": 10,
    "exposicao": "",
    "extracao": "RNeasy® Mini Kit (QIAGEN)",
    "quantificacao": "NanoDrop 2000 (Thermo Fisher Scientific)",
    "kit_cdna": "NZY First-Strand cDNA Synthesis Kit (NZYtech)",
    "massa_rna": 500,
    "termociclador": "Veriti (Applied Biosystems)",
    "master_mix": "qPCR Green Master LowROX with UNG (Cellco)",
    "equipamento_qpcr": "QuantStudio 3 (Thermo Fisher Scientific)",
    "volume_qpcr": 10,
    "ng_cdna": 10,
    "primer_um": 1.0,
}

TEST_EN = {"ANOVA de uma via": "one-way ANOVA", "ANOVA de Welch": "Welch's ANOVA",
           "Kruskal-Wallis": "Kruskal-Wallis test", "Teste t de Student": "Student's t-test",
           "Teste t de Welch": "Welch's t-test", "Mann-Whitney": "Mann-Whitney U test"}
POST_EN = {"Dunnett": "Dunnett's multiple comparisons test", "Tukey": "Tukey's multiple comparisons test",
           "Games-Howell": "the Games-Howell post hoc test", "Dunn (Holm)": "Dunn's test with Holm correction"}
POST_PT = {"Dunnett": "teste de comparações múltiplas de Dunnett", "Tukey": "teste de comparações múltiplas de Tukey",
           "Games-Howell": "teste post hoc de Games-Howell", "Dunn (Holm)": "teste de Dunn com correção de Holm"}
POST_SHORT = {"Dunnett": "Dunnett", "Tukey": "Tukey", "Games-Howell": "GH", "Dunn (Holm)": "Dunn", "—": ""}


def _n(x, lang, d=2):
    s = f"{x:.{d}f}"
    return s.replace(".", ",") if lang == "pt" else s


def _p(p, lang):
    if p is None or pd.isna(p):
        return "p = –"
    if p < 1e-4:
        s = "p < 0.0001"
    elif p < 0.001:
        s = f"p = {p:.4f}"
    elif p < 0.01:
        s = f"p = {p:.3f}"
    else:
        s = f"p = {p:.3f}" if p < 0.1 else f"p = {p:.2f}"
    return s.replace(".", ",") if lang == "pt" else s


def _stat(om, lang):
    v = _n(om["valor"], lang, 2)
    if om["estatística"] in ("F", "F de Welch"):
        d2 = om["df2"]
        d2s = f"{d2:.0f}" if float(d2).is_integer() else _n(d2, lang, 1)
        return f"F({om['df1']:.0f}, {d2s}) = {v}"
    if om["estatística"] == "H":
        return f"H({om['df1']:.0f}) = {v}"
    if om["estatística"] == "t":
        return f"t({_n(om['df1'], lang, 1)}) = {v}"
    return f"U = {v}"


def _fc_table(summary, gene, lvl):
    q = summary[summary["Gene"] == gene]
    if lvl and "Fator 2" in q:
        q = q[q["Fator 2"] == lvl]
    return dict(zip(q["Grupo"], q["Fold change médio"]))


def _join(items, lang):
    items = [items] if isinstance(items, str) else list(items)
    if len(items) == 1:
        return items[0]
    return ", ".join(items[:-1]) + (" and " if lang == "en" else " e ") + items[-1]


def _norm_sentence(refs, eff, lang):
    refs = [refs] if isinstance(refs, str) else list(refs)
    names = _join(refs, lang)
    if lang == "en":
        if len(refs) == 1 and not eff:
            return (f"The cycle threshold (Ct) of each target gene was normalized to the reference gene ({names}), and "
                    "relative expression was calculated using the 2^(−ΔΔCt) method.")
        how = "gene-specific amplification efficiencies and " if eff else ""
        norm = (f"the geometric mean of the reference genes ({names})" if len(refs) > 1
                else f"the reference gene ({names})")
        return ("Relative expression was calculated as calibrated normalized relative quantities (NRQ), using "
                f"{how}normalization to {norm}, according to the generalized model of Hellemans et al. (2007), with "
                "the mean of the control group as calibrator.")
    if len(refs) == 1 and not eff:
        return (f"O Ct de cada gene-alvo foi normalizado pelo gene de referência ({names}) e a expressão relativa foi "
                "calculada pelo método 2^(−ΔΔCt).")
    how = "a eficiência de amplificação de cada gene e " if eff else ""
    norm = (f"a média geométrica dos genes de referência ({names})" if len(refs) > 1
            else f"o gene de referência ({names})")
    return ("A expressão relativa foi calculada como quantidade relativa normalizada e calibrada (NRQ), considerando "
            f"{how}a normalização por {norm}, segundo o modelo generalizado de Hellemans et al. (2007), com a média "
            "do grupo controle como calibrador.")


def methods(cfg, ref_gene, n_por_grupo, lang, eff=False):
    c = {**METHODS_DEFAULTS, **cfg}
    tipo = _OLD.get(c["tipo_amostra"], c["tipo_amostra"])
    st_en, st_pt = SAMPLE_TYPES.get(tipo, (str(tipo).lower(), str(tipo).lower()))
    expo = f" ({c['exposicao']})" if c["exposicao"] else ""
    sp_en = c["especie"]
    sp_pt = c.get("especie_pt") or sp_en
    pool = int(c["animais_por_pool"] or 1)
    if lang == "en":
        of_sp = f" of {sp_en}" if sp_en else ""
        first = (f"Samples consisted of {n_por_grupo} pools per group, each containing {pool} {st_en}{of_sp}{expo}."
                 if pool > 1 else
                 f"Samples consisted of {n_por_grupo} biological replicates per group ({st_en}{of_sp}{expo}).")
        return (
            f"Relative gene expression was evaluated by qPCR. {first} Samples were stored in RNAlater "
            f"(Sigma-Aldrich®) at −20 °C until RNA extraction. Total RNA was purified using the {c['extracao']} "
            f"according to the manufacturer's instructions, and RNA concentration and purity (A260/280 and A260/230 "
            f"ratios) were assessed using a {c['quantificacao']}. cDNA was synthesized from {c['massa_rna']} ng of total "
            f"RNA using the {c['kit_cdna']} in a 20 µL reaction (25 °C for 10 min, 55 °C for 20 min and 85 °C for "
            f"5 min) in a {c['termociclador']} thermal cycler. qPCR reactions ({c['volume_qpcr']} µL) contained 1X "
            f"{c['master_mix']}, {_n(c['primer_um'], 'en', 1)} µM of each primer and cDNA equivalent to "
            f"{c['ng_cdna']} ng of RNA, and were run in technical duplicates on a {c['equipamento_qpcr']} system under "
            "the following conditions: 50 °C for 2 min (UNG), 95 °C for 10 min, and 40 cycles of 95 °C for 15 s and "
            "60 °C for 1 min, followed by a melt curve analysis to confirm amplification specificity. No-template "
            "and no-RNA (reverse transcription blank) controls were included for each gene. Primer efficiency was "
            "verified by standard curves (acceptance: 90–110%; R² ≥ 0.98). " + _norm_sentence(ref_gene, eff, "en"))
    return (
        "A expressão gênica relativa foi avaliada por qPCR. "
        + (f"As amostras consistiram em {n_por_grupo} pools por grupo, cada um com {pool} {st_pt}"
           f"{' de ' + sp_pt if sp_pt else ''}{expo}." if pool > 1 else
           f"As amostras consistiram em {n_por_grupo} réplicas biológicas por grupo ({st_pt}"
           f"{' de ' + sp_pt if sp_pt else ''}{expo}).")
        + " As amostras foram armazenadas em "
        f"RNAlater (Sigma-Aldrich®) a −20 °C até a extração. O RNA total foi purificado com o {c['extracao']}, "
        f"segundo as instruções do fabricante, e a concentração e a pureza do RNA (razões A260/280 e A260/230) foram "
        f"avaliadas no {c['quantificacao']}. O cDNA foi sintetizado a partir de {c['massa_rna']} ng de RNA total com o "
        f"{c['kit_cdna']} em reação de 20 µL (25 °C por 10 min, 55 °C por 20 min e 85 °C por 5 min) no termociclador "
        f"{c['termociclador']}. As reações de qPCR ({c['volume_qpcr']} µL) continham {c['master_mix']} 1X, "
        f"{_n(c['primer_um'], 'pt', 1)} µM de cada primer e cDNA equivalente a {c['ng_cdna']} ng de RNA, e foram "
        f"realizadas em duplicata técnica no {c['equipamento_qpcr']} nas condições: 50 °C por 2 min (UNG), 95 °C por "
        "10 min e 40 ciclos de 95 °C por 15 s e 60 °C por 1 min, seguidos de curva de dissociação (melting) para "
        "confirmar a especificidade. Controles sem molde (NTC) e brancos da transcrição reversa foram incluídos para "
        "cada gene. A eficiência dos primers foi verificada por curva padrão (aceite: 90–110%; R² ≥ 0,98). "
        + _norm_sentence(ref_gene, eff, "pt"))


TEST_PT = {"ANOVA de uma via": ("ANOVA de uma via", "seguida do"), "ANOVA de Welch": ("ANOVA de Welch", "seguida do"),
           "Kruskal-Wallis": ("teste de Kruskal-Wallis", "seguido do"),
           "Teste t de Student": ("teste t de Student", ""), "Teste t de Welch": ("teste t de Welch", ""),
           "Mann-Whitney": ("teste U de Mann-Whitney", "")}


def _combos(analyses):
    """{(teste, posthoc): [genes]} – para dizer qual gene usou qual teste."""
    out = {}
    for (gene, lvl), a in analyses.items():
        if "omnibus" in a:
            out.setdefault((a["teste"], a["posthoc"]), [])
            if gene not in out[(a["teste"], a["posthoc"])]:
                out[(a["teste"], a["posthoc"])].append(gene)
    return out


def _test_phrase(t, p, lang):
    if lang == "en":
        return TEST_EN.get(t, t) + (f" followed by {POST_EN.get(p, p)}" if p != "—" else "")
    name, link = TEST_PT.get(t, (t, "seguido do"))
    return name + (f", {link} {POST_PT.get(p, p)}" if p != "—" else "")


def statistics(analyses, metric, comparison, lang):
    combos = _combos(analyses)
    scale_en = {"ΔCt": "ΔCt values (equivalent to log2-transformed fold change)",
                "log2 fold change": "log2-transformed fold change values",
                "Fold change": "fold change (2^−ΔΔCt) values"}[metric]
    scale_pt = {"ΔCt": "valores de ΔCt (equivalentes ao fold change em log2)",
                "log2 fold change": "valores de fold change transformados em log2",
                "Fold change": "valores de fold change (2^−ΔΔCt)"}[metric]
    if lang == "en":
        parts = [f"Statistical analyses were performed on {scale_en} of each biological replicate (pool). Normality of "
                 "the residuals and homogeneity of variances were assessed with the Shapiro-Wilk and Brown-Forsythe "
                 "(Levene) tests, respectively, and the test for each gene was chosen accordingly."]
        if len(combos) == 1:
            (t, p), = combos
            parts.append(f"Groups were compared using {_test_phrase(t, p, 'en')}.")
        else:
            items = [f"{_test_phrase(t, p, 'en')} ({', '.join(g)})" for (t, p), g in combos.items()]
            parts.append("Groups were compared using " + "; ".join(items) + ".")
        parts.append("The significance level was set at α = 0.05. Analyses were performed in AmpliQA (Python; "
                     "SciPy and statsmodels).")
    else:
        parts = [f"As análises estatísticas foram feitas sobre os {scale_pt} de cada réplica biológica (pool). A "
                 "normalidade dos resíduos e a homogeneidade das variâncias foram avaliadas pelos testes de "
                 "Shapiro-Wilk e Brown-Forsythe (Levene), respectivamente, e o teste de cada gene foi escolhido de "
                 "acordo com essas premissas."]
        if len(combos) == 1:
            (t, p), = combos
            parts.append(f"Os grupos foram comparados por {_test_phrase(t, p, 'pt')}.")
        else:
            items = [f"{_test_phrase(t, p, 'pt')} ({', '.join(g)})" for (t, p), g in combos.items()]
            parts.append("Os grupos foram comparados por " + "; ".join(items) + ".")
        parts.append("O nível de significância adotado foi α = 0,05. As análises foram realizadas no AmpliQA "
                     "(Python; SciPy e statsmodels).")
    return " ".join(parts)


def results(analyses, summary, control, labels, lang, alpha=0.05):
    lab = lambda g: labels.get(g, g)
    out = []
    for (gene, lvl), a in analyses.items():
        if "omnibus" not in a:
            continue
        om = a["omnibus"]
        fc = _fc_table(summary, gene, lvl)
        at = (f" at {lvl}" if lang == "en" else f" em {lvl}") if lvl else ""
        stat = f"{_stat(om, lang)}, {_p(om['p'], lang)}"
        short = POST_SHORT.get(a["posthoc"], "")
        if om["p"] >= alpha:
            out.append(f"The expression of {gene}{at} did not differ among groups ({stat})." if lang == "en" else
                       f"A expressão de {gene}{at} não diferiu entre os grupos ({stat}).")
            continue
        sent = (f"The expression of {gene}{at} differed among groups ({stat})." if lang == "en" else
                f"A expressão de {gene}{at} diferiu entre os grupos ({stat}).")
        pv = a["pares"]
        ctrl_rows = pv[(pv["Grupo 1"] == control) | (pv["Grupo 2"] == control)]
        down, up, ns = [], [], []
        for _, r in ctrl_rows.iterrows():
            g = r["Grupo 2"] if r["Grupo 1"] == control else r["Grupo 1"]
            f = fc.get(g, np.nan)
            ps = _p(r["p ajustado"], lang).replace("p ", f"p{short} " if short else "p ", 1)
            fcs = f"{_n(f, lang)}-fold change" if lang == "en" else f"fold change = {_n(f, lang)}"
            item = f"{lab(g)} ({fcs}, {ps})"
            (ns if r["p ajustado"] >= alpha else down if f < 1 else up).append(item)

        def join(items):
            if len(items) == 1:
                return items[0]
            sep = " and " if lang == "en" else " e "
            return ", ".join(items[:-1]) + sep + items[-1]

        post = a["posthoc"]
        parts = []
        if lang == "en":
            if down:
                parts.append(f"downregulation in {join(down)}")
            if up:
                parts.append(f"upregulation in {join(up)}")
            if parts:
                sent += f" {POST_EN.get(post, post).capitalize()} revealed " + "; and ".join(parts) + "."
            if ns:
                sent += f" No difference from the control was found in {join(ns)}."
        else:
            if down:
                parts.append(f"redução da expressão em {join(down)}")
            if up:
                parts.append(f"aumento da expressão em {join(up)}")
            if parts:
                sent += f" O {POST_PT.get(post, post)} revelou " + "; e ".join(parts) + "."
            if ns:
                sent += f" Não houve diferença em relação ao controle em {join(ns)}."
        others = pv[(pv["Grupo 1"] != control) & (pv["Grupo 2"] != control) & (pv["p ajustado"] < alpha)]
        for _, r in others.iterrows():
            ps = _p(r["p ajustado"], lang)
            sent += (f" {lab(r['Grupo 1'])} and {lab(r['Grupo 2'])} also differed from each other ({ps})."
                     if lang == "en" else
                     f" {lab(r['Grupo 1'])} e {lab(r['Grupo 2'])} também diferiram entre si ({ps}).")
        out.append(sent)
    return "\n\n".join(out)


def legend(genes, ref_gene, n_por_grupo, err, analyses, lang):
    combos = _combos(analyses)
    refs = [ref_gene] if isinstance(ref_gene, str) else list(ref_gene)
    calc = "NRQ" if len(refs) > 1 else "2^−ΔΔCt"
    if len(refs) > 1:
        ref_gene = ('the geometric mean of ' if lang == 'en' else 'média geométrica de ') + _join(refs, lang)
    else:
        ref_gene = refs[0]
    thr = "* p < 0.05; ** p < 0.01; *** p < 0.001; **** p < 0.0001"
    if len(combos) == 1:
        t = _test_phrase(*next(iter(combos)), lang)
    else:
        t = "; ".join(f"{_test_phrase(a, b, lang)} ({', '.join(g)})" for (a, b), g in combos.items())
    if lang == "en":
        e = "SEM" if err == "SEM" else "SD"
        return (f"Figure X – Relative gene expression (fold change, {calc}) of {', '.join(genes)}, normalized to "
                f"{ref_gene} (n = {n_por_grupo} pools per group). Bars represent mean ± {e}; dots represent "
                f"individual pools. Statistics: {t}. {thr}.")
    e = "erro padrão da média (EPM)" if err == "SEM" else "desvio padrão (DP)"
    return (f"Figura X – Expressão gênica relativa (fold change, {calc}) de {', '.join(genes)}, normalizada por "
            f"{ref_gene} (n = {n_por_grupo} pools por grupo). As barras representam a média ± {e}; os pontos são os "
            f"pools individuais. Estatística: {t}. {thr.replace('0.0', '0,0')}.")


def to_docx(sections: dict, title: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Roboto"
    st.font.size = Pt(11)
    doc.add_heading(title, level=1)
    for head, body in sections.items():
        doc.add_heading(head, level=2)
        for para in str(body).split("\n\n"):
            doc.add_paragraph(para)
    doc.add_paragraph("Rascunho gerado pelo AmpliQA – revise antes de usar.").runs[0].italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
